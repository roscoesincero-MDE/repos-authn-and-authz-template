"""Add G-22 and G-23, found while building Phases 3 and 4, and refresh the Legend counts."""
import copy
from openpyxl import load_workbook

WB = 'workbooks/gaps.xlsx'

ROWS = [
    ('G-22',
     'The authorization hot path is the only unmeasured path in the database',
     'DES section 20.2 (performance targets); section 21.2 (logging)',
     'Section 20.2 sets targets for the authorization path -- a permission check and a filtered read must not be the '
     'reason a screen is slow -- and section 21.2 expects every module to be observable through logs.ExecutionLog.',
     'Neither, for the hot path, and not by oversight. auth.udfHasPermission, auth.tvfPermissionScope and the three '
     'predicates are functions: a function cannot write to a table, and a SCHEMABINDING predicate must not try. The '
     'predicates run once per row per statement, so per-call logging would cost more than the check. So the one part '
     'of this database that runs on every single query is the one part that reports nothing about itself, while the '
     'procedures around it all log.',
     'The cost is diagnostic, not functional, and it lands on the day a deployment is slow. Today the only way to '
     'measure the predicate is Query Store or an extended-events session set up by hand by whoever suspects it, and '
     'the answer they get is aggregated across every tenant and every profile -- which is the wrong shape, because the '
     'interesting case is one profile with authority at many tenants, or one tenant with a deep subtree. There is also '
     'no baseline: nobody will know whether 40 ms is normal for this workload or a regression introduced by seeding '
     'the catalogue, adding a permission to the literal list, or registering a third table. A gap that costs nothing '
     'until it costs a week.',
     'Medium',
     'Phase 5 performance work (T-069 onwards); any capacity conversation; diagnosing a slow screen in production',
     'A SAMPLED probe rather than per-call logging. Add a config.ApplicationSetting row, Authz.InstrumentationSampleRate, '
     'defaulting to 0. Where a procedure calls auth.uspDemandPermission it may -- when the sample fires -- record the '
     'elapsed time, the profile, the acting tenant, the permission and the row count into a new logs table, which the '
     'PROCEDURE writes rather than the function. For the predicates themselves, ship a documented extended-events '
     'session definition and a query over sys.dm_exec_function_stats in the Phase 5 test file, so that measuring is a '
     'documented step and not an investigation. Do not instrument the functions.',
     None,
     '5 Performance',
     'Open',
     'Found at the end of Phase 4, while writing the closing report of 120_rls_policy.sql: the report can say how many '
     'tables are bound and which permission ids are in the predicate, and cannot say what any of it COSTS. Recorded '
     'now rather than in Phase 5 because the shape of the resolution constrains Phase 5: if the sampling hook is not '
     'designed alongside the indexes, the measurement will be bolted on afterwards by whoever is already in a hurry.'),

    ('G-23',
     'A new table in the logs schema is writable by applicationRole until someone remembers to deny it',
     'DES section 13; 085_logs_auth_tables.sql section 8; conventions scripts/permissions.sql',
     'The audit trails are written only through the recorder procedures in 165_logs_procedures.sql, by ownership '
     'chaining. No role holds table-level write access to a trail, so a compromised application login cannot forge or '
     'erase a record of what it did.',
     'True today, and true by a list rather than by a rule. The conventions skill\'s scripts/permissions.sql grants '
     'SELECT, INSERT, UPDATE and DELETE on SCHEMA::logs to applicationRole, so every table in that schema is writable '
     'by default; 170_permissions.sql now names the four trail tables and DENIES the three write verbs on each, and '
     'asserts all twelve denies. The rule is inherited and permissive; the exception is enumerated and restrictive.',
     'The next logs table repeats the hole, silently and by default. Whoever adds logs.ConsentEvent in Phase 6 gets a '
     'table the application can INSERT into, UPDATE and DELETE from directly, and nothing fails: the closing report of '
     '170 checks the four tables it knows about, so it will still report 12 of 12 OK. That is the worst available '
     'shape for a control -- a green report next to an ungoverned table. The severity is Medium rather than High only '
     'because ownership chaining means no legitimate caller needs the grant, so removing it breaks nothing; the risk '
     'is entirely that nobody notices it is there.',
     'Medium',
     'Any future logs table; T-089 and Phase 6 logging work; the 950 verification report (T-101)',
     'Invert the check. Instead of listing the tables that must be denied, have 170_permissions.sql enumerate '
     'sys.tables in the logs schema, subtract the documented exception (logs.ExecutionLog, which the framework writes '
     'from the application side), and assert that applicationRole holds a write DENY on every remaining table -- '
     'failing the deployment on any table it has not seen before. That turns "add a table, remember a DENY" into "add '
     'a table, the deployment tells you". The same inversion belongs in 950_verify_deployment.sql so it is checked '
     'between deployments too.',
     None,
     '8 Verification',
     'Open',
     'Found during the Phase 4 closeout by a catalog query, not by reading any file: 085 section 8 said nothing is '
     'granted on the trail tables, the conventions script had granted everything on the schema, and both statements '
     'were true. BL-048 records the narrowing that closed the immediate hole and the NOTE row in 170 that warns about '
     'the next table; this row is the standing gap, because a warning in a report is not a control.'),
]


def main():
    wb = load_workbook(WB)
    ws = wb['Gaps']
    tpl = ws.max_row                      # G-21
    at = tpl + 1
    for i, vals in enumerate(ROWS):
        rn = at + i
        for j, v in enumerate(vals):
            ws.cell(row=rn, column=j + 1).value = v
        for c in range(1, 14):
            ws.cell(row=rn, column=c)._style = copy.copy(ws.cell(row=tpl, column=c)._style)
        ws.row_dimensions[rn].height = ws.row_dimensions[tpl].height

    if ws.auto_filter.ref:
        head, tail = ws.auto_filter.ref.split(':')
        lastcol = ''.join(ch for ch in tail if ch.isalpha())
        ws.auto_filter.ref = '%s:%s%d' % (head, lastcol, at + len(ROWS) - 1)

    lg = wb['Legend']
    lg.cell(row=24, column=2).value = (
        "Twenty-three gaps as of 2026-09-20, against DES-AUTH-001 v1.5. Three are CLOSED: G-19 by "
        "170_permissions.sql, G-20 before Phase 3 began, and G-07 by a management decision on 2026-09-20 -- each with "
        "its resolution recorded in the Notes column rather than deleted. G-07 is the gap that actually cost this "
        "project something while it was open: it blocked T-041, so MFA enrolment did not exist, everything that "
        "CONSUMES a second factor had to be tested against a factor inserted directly by db_owner, and one fixture "
        "account was permanently unable to sign in. All of that is now resolved, and the db_owner insert was KEPT "
        "deliberately, as the control that proves the consuming experiments do not depend on the enrolling ones.\n\n"
        "HOW THE LAST THREE WERE FOUND is the part worth copying. G-21 came from the Phase 2 test being unable to "
        "write an assertion. G-22 came from writing a closing report that could say how many tables were protected and "
        "not what the protection cost. G-23 came from a catalog query during the Phase 4 closeout, where two "
        "statements were both true -- 085 said nothing is granted on the trail tables, the conventions script had "
        "granted everything on the schema. None of the three was found by re-reading the design. Add rows as the build "
        "finds more; a gap discovered during implementation is the normal case, not a failure of the design pass."
    )
    wb.save(WB)
    print('gaps added', len(ROWS), 'filter', ws.auto_filter.ref)


if __name__ == '__main__':
    main()
