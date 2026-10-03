# Records the fifth template defect found by the WellDrillersLicense build and fixed here on 2026-09-25:
# G-56 in gaps.xlsx and BL-087 in the build log. No new script, so the Scripts sheet is unchanged.
import copy
from openpyxl import load_workbook

GAPS = r'../workbooks/gaps.xlsx'
BUILD = r'../workbooks/build-and-traceability.xlsx'

G = [
    ['G-56', 'Re-running 115 withdrew every UiElementPermission gate it did not list, so an element a consumer added lost its View gate and became visible to every profile',
     'DES 13.1, 13.2; T-090; G-18',
     'A consumer extends the starter catalogue (T-090) and re-runs 115 safely: 115 converges its own elements and gates and leaves everyone else\'s alone.',
     'The gate MERGE\'s WHEN NOT MATCHED BY SOURCE covered the whole application\'s gates, not the elements 115 seeds. The CaseFile demo elements were also seeded whether or not the demo domain (dbo.CaseFile) was installed.',
     'Fail-open. An element with no View gate is visible to every authenticated profile (section 13.2), so one re-run of 115 opened a consumer\'s whole menu, with CanEdit off. A consumer without the demo domain got a Cases menu gated on Data.Read, which every staff role holds.',
     'High', 'Any consumer that extends the catalogue, or installs without the demo domain',
     'Gate withdrawal is scoped to the elements 115 lists. The demo elements (Area.Cases and its descendants, taken from the tree) are seeded only where dbo.CaseFile exists, and soft-deleted where it does not. The closing report expects 26 starter elements with the demo and 12 without.',
     'None', '3 Authorization', 'Closed',
     'CLOSED 2026-09-25 in 115_seed_reference_data.sql. Where the demo is installed, as in every template build, the seeded catalogue is unchanged. Found by the WellDrillersLicense build (its GAP-WDL-021 / BL-WDL-069). The catalogue section is identical in both copies. It was proved there on a fresh build without the demo (12 starter elements; 115 re-run, the 136 WDL gates kept) and on one with it (26 of 26; 010-080 pass). tplUpstream was not rebuilt for it.'],
]

BL = ['BL-087', '2026-09-25',
      'database/115_seed_reference_data.sql',
      'DES 13.1, 13.2; T-090; G-56',
      'The starter catalogue is a starting point that consumers extend (T-090); an element with no View gate is visible to everyone.',
      'The gate MERGE withdraws only gates on the elements 115 seeds. The CaseFile demo elements are seeded only where dbo.CaseFile exists. The closing report\'s expected count follows (26 or 12).',
      'Nothing in the design changes: the fix makes 115 do what T-090 already assumed. With the demo installed, 115 seeds exactly what it did before.',
      'Build, carrying back a downstream finding.',
      'No design change needed.']


def append(ws, rows):
    tpl = ws.max_row
    for row in rows:
        rn = ws.max_row + 1
        for c, v in enumerate(row, start=1):
            ws.cell(row=rn, column=c, value=v)
            ws.cell(row=rn, column=c)._style = copy.copy(ws.cell(row=tpl, column=c)._style)


wb = load_workbook(GAPS)
ws = wb['Gaps']
assert not any(r[0] == 'G-56' for r in ws.iter_rows(min_row=2, values_only=True)), 'already recorded'
append(ws, G)
wb.save(GAPS)

wb = load_workbook(BUILD)
ws = wb['Build Log']
assert not any(r[0] == 'BL-087' for r in ws.iter_rows(min_row=2, values_only=True)), 'already recorded'
append(ws, [BL])
wb.save(BUILD)
print('workbooks updated')
