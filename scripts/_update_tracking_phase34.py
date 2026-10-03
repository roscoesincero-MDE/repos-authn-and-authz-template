"""Mark Phase 3 (T-043..T-058) and Phase 4 (T-059..T-068) complete in implementation-tracking.xlsx.

T-045 does not exist -- the ID is skipped in the Tasks sheet and always has been.
"""
import copy
import datetime as dt
from openpyxl import load_workbook

WB = 'workbooks/implementation-tracking.xlsx'
D = dt.datetime(2026, 9, 20)

# task id -> (actual days, note)
UPDATES = {
    'T-043': (0.8,
        'auth.PermissionCategory and auth.Permission built in database/050_auth_permission.sql, with IsTenantScoped '
        'NOT NULL on Permission and the category/code pair kept honest by FK_auth_Permission_PermissionCategory over '
        '(PermissionCategoryId, PermissionCategoryCode). THE 35 ROWS THEMSELVES ARE NOT HERE: seeding the catalogue is '
        'T-089 (115_seed_reference_data.sql, Phase 6), which is what the Task column means by "the full catalogue" and '
        'is where IsTenantScoped per permission is actually set. What this task delivered is the structure that makes '
        'the catalogue expressible, plus PermissionDescription -- an addition to DES 15.4 recorded as BL-037, because '
        'the alternative is 35 meanings hard-coded in the UI project. Deployed and verified on '
        'MDE-55TT2J4\\testTemplate; the empty catalogue has a visible consequence recorded against T-065.'),
    'T-044': (0.8,
        'auth.Role and auth.RolePermission built in database/055_auth_role.sql. UX_auth_Role_Grant is filtered UNIQUE '
        'on (ApplicationId, OwnerTenantId, RoleCode) where IsDeleted = 0, IsSystemRole is present and defended, and '
        'INV-04 (a role and the profile it is granted to belong to one application) is enforced structurally rather '
        'than by a procedure check: FK_auth_Role_OwnerTenant references the unfiltered UX_auth_Tenant_Id_Application '
        'pair added to 030_auth_tenant.sql by ALTER for this purpose -- BL-040. The fourteen baseline roles are seed '
        'data and belong to T-089.'),
    'T-046': (0.9,
        'auth.UserProfile built in database/040_auth_userprofile.sql, alongside auth.User rather than in an '
        'authorization script, and the banner argues the split. UX_auth_UserProfile_Id_Tenant is the UNFILTERED unique '
        'pair on (UserProfileId, TenantId) that the demo domain\'s composite foreign keys reference -- which is what '
        'stops a case file being assigned to a profile at another tenant, and what _tests/060 section 6g ran into as '
        'Msg 547 before the RLS block predicate was ever consulted. One default profile per person is '
        'UX_auth_UserProfile_Default, filtered on IsDefault = 1 AND IsDeleted = 0 (INV-03). ProfileName is unique per '
        '(UserId, TenantId) among live rows, an addition to DES 15.4 recorded as BL-036.'),
    'T-047': (0.7,
        'auth.UserProfileRole built in database/060_auth_profile_role.sql -- the only table in this database that '
        'records a human decision. ScopeTenantId, GrantedByProfileId, GrantedUtc and ExpiresUtc are all present; '
        'CK_auth_UserProfileRole_Expiry requires ExpiresUtc > GrantedUtc, which _tests/050 had to respect by '
        'backdating GrantedUtc to express an already-expired grant. UX_auth_UserProfileRole_Grant is filtered UNIQUE '
        'on (UserProfileId, RoleId, ScopeTenantId) where IsDeleted = 0 -- the reason a soft-deleted grant cannot be '
        'restored wholesale (Msg 2601, found by _tests/050 step 6).'),
    'T-048': (0.6,
        'auth.ProfilePermissionScope built in database/065_auth_effective_permission.sql: PK on all three columns '
        '(UserProfileId, PermissionId, ScopeTenantId), plus IX_auth_ProfilePermissionScope_Lookup filtered on '
        'IsDeleted = 0 so the RLS predicates can seek it. The filter is load-bearing and is BL-039: every reader MUST '
        'carry IsDeleted = 0 in its own WHERE clause or it both misses the index and counts retired rows as live '
        'authority. Derived table, never edited by hand.'),
    'T-049': (1.2,
        'auth.uspRebuildProfilePermissionScope built in database/065_auth_effective_permission.sql, taking one profile '
        'or the whole database, and excluding expired grants, deleted rows at every level and inactive or deleted '
        'users and profiles. Verified by _tests/050, which is the real exit criterion for this task: fourteen '
        'mutations -- grants, revocations, re-grants, role edits, expiries, a profile deactivation and a user '
        'deactivation -- and after every one of them the materialized table equalled the set derived longhand from DES '
        '8.6 in BOTH directions, 14 exercised, 0 failed. The concurrency half of gap G-03 remains open and is recorded '
        'there, not here.'),
    'T-050': (0.7,
        'auth.udfHasPermission built in database/100_auth_functions.sql: two seeks and no join to auth.Permission by '
        'code, because the code-to-id resolution is the caller\'s or the deploy step\'s job. Platform-category '
        'permissions additionally require IsPlatformAdmin on the user (INV-09), which is the one asymmetry in the '
        'function and is commented as such. Exercised through auth.uspDemandPermission by _tests/050 and by the '
        'deployment probes in 150.'),
    'T-051': (0.5,
        'auth.tvfPermissionScope built in database/100_auth_functions.sql as an INLINE table-valued function -- tvf, '
        'not udf, and the rename is BL-041: the house convention reserves udf for scalars, the SQL gate enforces it, '
        'and the design document was corrected rather than the code. Returns the tenants a profile holds one '
        'permission on. WHAT IT DOES NOT DO is filter unusable tenants (UI-33): it answers "where is this authority '
        'recorded", not "where can it be exercised today", and a caller that needs the second question must join '
        'auth.udfIsTenantUsable itself.'),
    'T-052': (0.6,
        'auth.uspDemandPermission built in database/150_auth_query_procedures.sql, raising E-50030 when the profile '
        'lacks the permission anywhere and E-50031 when it holds it but not at the acting tenant -- two numbers '
        'because the UI response differs (UI-13). Writes logs.AuthorizationDenial through '
        'logs.uspRecordAuthorizationDenial before it throws, which is why 165 installs before 150. THE LIMITATION IS '
        'STATED, NOT PAPERED OVER: called inside a caller\'s open transaction, the rollback the denial provokes takes '
        'the trail row with it (BL-042), so DES section 9 authorizes at step 5, before any transaction opens, and '
        'every procedure in this database follows that order.'),
    'T-053': (1.5,
        'auth.uspSetSessionContext built in database/105_auth_session_procedures.sql: the full validation chain '
        '(session hash, expiry, user usable, profile usable, tenant usable, application match) and the three-way '
        'already-set branch -- same profile returns silently, different profile raises E-50022, no profile set '
        'proceeds. It takes @SessionTokenHash and NOT the token; the design document\'s sketch of @SessionToken was '
        'corrected instead of the code, BL-043, because the token never crosses this boundary. On an expired session '
        'it calls auth.uspEndSession before throwing E-50023, which is why the installer runs 105 after 110.'),
    'T-054': (0.2,
        'auth.uspClearSessionContext built in database/105_auth_session_procedures.sql. It clears what CAN be cleared '
        'and says so: the five identity keys are set @read_only = 1 and cannot be re-set or nulled for the life of the '
        'connection (Msg 15664), so this procedure clears BypassRowSecurity and reports the rest as a connection-pool '
        'fact rather than pretending to reset them. That asymmetry is UI-06 and is why one connection serves exactly '
        'one profile.'),
    'T-055': (0.7,
        'logs.AuthorizationChange, logs.AuthorizationDenial and logs.DataChangeLog built in '
        'database/085_logs_auth_tables.sql, in logs rather than logsData because a human reads them -- which is what '
        'lets logsAuditReader read all three directly without anything on SCHEMA::auth. ActorAuthorityTenantId is '
        'present on the authority trail (P-08). Closed during the Phase 3/4 closeout: the conventions\' own '
        'scripts/permissions.sql grants all four table verbs on SCHEMA::logs to applicationRole, which covered these '
        'three the moment they were created there, so 170_permissions.sql now DENIES INSERT, UPDATE and DELETE on '
        'them object by object -- BL-048.'),
    'T-056': (0.8,
        'logs.uspRecordAuthorizationChange, logs.uspRecordAuthorizationDenial and logs.uspRecordDataChange built in '
        'database/165_logs_procedures.sql, all three carrying ActorAuthorityTenantId. EXECUTE is granted to NOBODY, '
        'deliberately: the callers reach them by ownership chaining -- measured on this instance, and re-measured '
        'during the closeout with the new object-level DENY in place, where a direct UPDATE by an applicationRole '
        'member failed with 229 while the same write through auth.uspRecordLoginFailure succeeded. 170_permissions.sql '
        'now asserts that absence as well as 165 asserting it.'),
    'T-057': (1.4,
        'database/_tests/050_authorization_and_session.sql, sections 3 and 4. The arbitrary sequence is expressed as '
        'DATA -- a 14-row @Mutation table driven by one WHILE loop -- so that the derived-set definition from DES 8.6 '
        'exists exactly once and is written longhand rather than reusing the procedure\'s own query, which would only '
        'have proved self-consistency. Both EXCEPT directions are checked and labelled by consequence: missing rows '
        'are a FALSE DENIAL (the user sees an empty screen), extra rows are a FALSE GRANT (nobody complains). 14 '
        'exercised, 0 failed; the file is re-runnable and was run twice. Three real constraints were found the hard '
        'way and are written into the expectations: CK_auth_UserProfileRole_Expiry, UX_auth_UserProfileRole_Grant '
        '(Msg 2601 on restoring every soft-deleted grant), and trg_au_updt_UserProfileRole overwriting auditDeletedBy '
        'so it cannot carry a caller\'s marker (BL-044).'),
    'T-058': (0.4,
        'database/_tests/050_authorization_and_session.sql section 5, which MUST be the last section in the file: '
        'setting the identity keys spends the connection, so the section that tests the already-set branch cannot be '
        'followed by anything needing a different profile. Recorded: the same profile twice returns silently with '
        'UserProfileId intact, a different profile raises Error 50022 (expected 50022) and the identity keys survive '
        'the refusal -- @IdentityKeysRemain = 1. Section 4b therefore sets UserProfileId and ActingTenantId directly, '
        'writable then nulled, to get a profile identity without spending the connection.'),
    'T-059': (0.4,
        'config.ApplicationSetting built in database/025_config_tables.sql and seeded with 18 live rows, verified on '
        'MDE-55TT2J4\\testTemplate. It installs seventh, before any auth table, because 110\'s authentication '
        'procedures read four settings out of it -- install order is not build order, and this row is why the manifest '
        'comment says so. IsSensitive is present and carries weight: Authn.DummyVerifierPepper is why '
        '170_permissions.sql DENIES SELECT on the whole table to both reading roles, with a filtered view named as the '
        'extension point when something needs a direct read.'),
    'T-060': (0.4,
        'config.TenantScopedTable built in database/025_config_tables.sql with 2 live rows -- dbo.CaseFile and '
        'dbo.CaseNote -- and it is the registry auth.uspRebuildTenantAccessPolicy reads to decide what to protect. A '
        'row here is a claim that the table carries a TenantId column of the right type; the rebuild procedure '
        'verifies that claim and reports skipped rows through @TablesSkipped rather than binding a predicate to a '
        'column that is not there.'),
    'T-061': (1.0,
        'auth.tvfTenantReadPredicate built in database/100_auth_functions.sql: WITH SCHEMABINDING, TRY_CAST over every '
        'SESSION_CONTEXT read (a predicate that throws on a malformed key would fail the query rather than deny the '
        'row), the permission id list substituted as a LITERAL at deploy time, and the BypassRowSecurity key as the '
        'one escape. Named tvf and not udf per the house convention -- BL-041, and the design document was corrected. '
        'Verified by _tests/060 sections 6a and 6h.'),
    'T-062': (0.7,
        'auth.tvfTenantInsertPredicate built in database/100_auth_functions.sql: the row\'s tenant must EQUAL the '
        'acting tenant and the profile\'s insert scope must cover it (P-06). Equality and not descendant-of, which is '
        'the whole point: a manager with authority over a subtree still inserts into the tenant they are acting for, '
        'and the UI never offers a TenantId picker on an insert (UI-14). Verified by _tests/060 sections 6b and 6c -- '
        'the refusal is Msg 33504.'),
    'T-063': (0.5,
        'auth.tvfTenantUpdatePredicate built in database/100_auth_functions.sql: the profile\'s update scope must '
        'cover the row\'s tenant. It is bound three times per protected table -- BLOCK AFTER INSERT, BLOCK BEFORE '
        'UPDATE and BLOCK AFTER UPDATE -- and it is therefore also the predicate that governs the soft delete, since '
        'deletion in this design is an UPDATE. There is deliberately no delete predicate. Verified by _tests/060 '
        'sections 6d to 6g.'),
    'T-064': (1.3,
        'auth.TenantAccessPolicy built in database/120_rls_policy.sql from config.TenantScopedTable: FILTER plus three '
        'BLOCK predicates per registered table, 2 tables bound and 8 predicates, state ON, verified on '
        'MDE-55TT2J4\\testTemplate. The policy is built by auth.uspRebuildTenantAccessPolicy @Action = '
        'N\'Rebuild\'|N\'Drop\' rather than by inline DDL, because a bound function cannot be altered (error 3729) and '
        'a re-deployment therefore needs a supported way to unbind first -- BL-038. '
        'Install-TemplateDatabase.ps1 now calls @Action = N\'Drop\' before every pass and rebuilds at step 27.'),
    'T-065': (0.6,
        'The permission-id substitution and the assertion that keeps it honest, both in database/120_rls_policy.sql. '
        'The predicate carries a literal PermissionId IN (...) list because a database serving four applications holds '
        'four rows whose PermissionCode is Data.Read, and one literal would protect one application and silently fail '
        'open for the other three -- BL-039. The closing report re-derives the list from the catalogue and compares it '
        'with CHARINDEX against the deployed definition, so a stale list is a failed deployment rather than a silent '
        'denial (UI-35). With the catalogue unseeded until T-089 the list resolves to the sentinel -1 and every '
        'predicate denies every non-maintenance session: fail-CLOSED, reported by the script, and the reason '
        '115_seed_reference_data.sql must re-run this file.'),
    'T-066': (1.2,
        'auth.uspBeginMaintenanceSession and auth.uspEndMaintenanceSession built in '
        'database/105_auth_session_procedures.sql: gated on IS_ROLEMEMBER(\'rlsBypassRole\') with E-50070 on refusal, '
        'reason and ticket reference required, and the logs.AuthenticationEvent row written BEFORE the key is set so '
        'that a bypass cannot exist without a trail. ORIGINAL_LOGIN() is recorded, so the trail names the real login '
        'even under EXECUTE AS. A REAL DEFECT WAS FOUND AND FIXED HERE by _tests/060: both procedures wrote UserId and '
        'UserName NULL, which CK_logs_AuthenticationEvent_Attributable refuses (Msg 547) -- the deployment probe only '
        'ever reached the refusal path, so nothing had exercised the accepted path. Both now write the actor into '
        'UserName. BL-046.'),
    'T-067': (0.9,
        'database/_tests/060_row_security.sql sections 6a to 6h -- the full DES 10.3 matrix over four case files and '
        'two notes in four tenants: read (3 of 4 visible), insert into the acting tenant (1 row), insert into a '
        'sibling (33504), insert into a descendant (33504), update inside scope (1 row), update a row the profile can '
        'read but not write (33504), the tenant move (33504) and the soft delete (1 row). All eight cells came out as '
        'the design says they should, inside a run that recorded 16 observations as intended with 3 notes. Two '
        'measured facts are written into the file: a block violation is Msg 33504, and the composite foreign keys on '
        '(AssignedToProfileId, TenantId) and (CaseFileId, TenantId) refuse a tenant move with Msg 547 BEFORE the block '
        'predicate is consulted -- so 6b inserts an unassigned row for 6g to move. BL-045.'),
    'T-068': (0.3,
        'database/_tests/060_row_security.sql section 5 and section 7. The demonstration is a set-key / count / '
        'clear-key / count sequence so the bypass key is provably the only variable: with it set the connection sees '
        '4 of 4 case files, cleared it sees 0, and it is a db_owner connection either way -- RLS applies to db_owner. '
        'Section 7 then does it the supported way, through a user created WITHOUT LOGIN in rlsBypassRole calling '
        'auth.uspBeginMaintenanceSession: 4 rows inside the window, 0 after auth.uspEndMaintenanceSession, '
        '@BypassCleared = 1, and one MaintenanceBypass plus one MaintenanceBypassEnded row naming the real login. This '
        'is UI-18, and the point of writing it down is that an empty grid in SSMS looks exactly like data loss.'),
}


def main():
    wb = load_workbook(WB)
    ws = wb['Tasks']
    style_src = None

    for row in ws.iter_rows(min_row=2):
        tid = row[0].value
        if tid == 'T-042':
            style_src = row
        if tid not in UPDATES:
            continue
        actual, note = UPDATES[tid]
        ws.cell(row=row[0].row, column=10).value = 'Complete'      # Status
        ws.cell(row=row[0].row, column=11).value = 1               # % Done
        ws.cell(row=row[0].row, column=12).value = D               # Started
        ws.cell(row=row[0].row, column=13).value = D               # Completed
        ws.cell(row=row[0].row, column=14).value = actual          # Actual Days
        ws.cell(row=row[0].row, column=15).value = note            # Notes / Blockers

        if style_src is not None:
            for col in (10, 11, 12, 13, 14, 15):
                tgt = ws.cell(row=row[0].row, column=col)
                tgt._style = copy.copy(style_src[col - 1]._style)

    wb.save(WB)
    print('Tasks updated:', len(UPDATES))


if __name__ == '__main__':
    main()
