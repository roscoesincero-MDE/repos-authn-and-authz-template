"""Update the Scripts sheet of build-and-traceability.xlsx for Phases 3 and 4.

Rows that were 'Not started' predictions are rewritten with what actually shipped; rows that were already
Verified get an appended paragraph rather than a replacement, so the earlier phases' record survives.
Two rows are added for the new test files.
"""
import copy
from openpyxl import load_workbook

WB = 'workbooks/build-and-traceability.xlsx'
DEV = '2026-09-20  MDE-55TT2J4\\testTemplate'

# script name -> {column letter: new value}, notes appended when key is 'H+'
APPEND = {
    '025_config_tables.sql': {
        'H+': " BUILT AND CLOSED IN PHASE 4 (T-059, T-060): 18 live settings and 2 registry rows on dev. "
              "config.TenantScopedTable is what auth.uspRebuildTenantAccessPolicy reads to decide what to protect, and "
              "a row in it is a CLAIM that the named table carries a TenantId column of the right type -- the rebuild "
              "verifies the claim and reports failures through @TablesSkipped rather than binding a predicate to a "
              "column that is not there.",
        'J': DEV,
        'K': 'Opus 5 (T-025 dependency; T-059, T-060)',
    },
    '030_auth_tenant.sql': {
        'H+': " AMENDED IN PHASE 3 (BL-040): UX_auth_Tenant_Id_Application, an UNFILTERED UNIQUE (TenantId, "
              "ApplicationId), is added here by a guarded ALTER because auth.Role's composite foreign key needs it to "
              "make INV-04 structural rather than a check inside a procedure. It is an ALTER and not a line in the "
              "CREATE TABLE so that a database deployed before Phase 3 gains it on a re-run; the closing report names "
              "it either way.",
        'J': DEV,
        'K': 'Opus 5 (T-014, T-015, T-016, T-017, T-024; amended for T-044)',
    },
    '040_auth_userprofile.sql': {
        'D': 'auth.User, auth.UserProfile',
        'H+': " NO LONGER PARTIAL, as of Phase 3: auth.UserProfile is now in this file (T-046), which is where the "
              "banner always said it would go. It carries UX_auth_UserProfile_Id_Tenant -- the UNFILTERED unique "
              "(UserProfileId, TenantId) that the demo domain's composite foreign keys point at, and the reason a "
              "tenant move is refused with Msg 547 before any block predicate is consulted -- one default profile per "
              "person (UX_auth_UserProfile_Default, INV-03) and ProfileName unique per user and tenant among live rows "
              "(BL-036). The runner's RequiresAuth probe and the named skip of 090 are deliberately kept: both remain "
              "meaningful for a partial deployment.",
        'J': DEV,
        'K': 'Opus 5 (T-025, T-046)',
    },
    '085_logs_auth_tables.sql': {
        'H+': " THE THREE AUTHORIZATION TRAILS WERE BUILT IN PHASE 3 (T-055), in logs and not logsData because a human "
              "reads them -- which is what lets logsAuditReader read all three with nothing granted on SCHEMA::auth. "
              "Two things were then found from outside this file. A defect: CK_logs_AuthenticationEvent_Attributable "
              "requires UserId or UserName and the maintenance procedures in 105 were writing both NULL, caught by "
              "_tests/060 because the deployment probe had only ever reached the refusal path (BL-046). And a hole: the "
              "conventions' own permission script grants all four table verbs on SCHEMA::logs to applicationRole, so "
              "these tables were writable directly the moment they were created here -- 170_permissions.sql now denies "
              "INSERT, UPDATE and DELETE on each of them (BL-048).",
        'J': DEV,
        'K': 'Opus 5 (T-033, T-055)',
    },
    '100_auth_functions.sql': {
        'D': 'auth.udfIsTenantUsable, auth.udfIsUserUsable, auth.udfResolveAuthPolicy, auth.udfResolveSessionUser, '
             'auth.udfResolveEnrolmentActor, auth.udfHasPermission, auth.tvfPermissionScope, '
             'auth.tvfTenantReadPredicate, auth.tvfTenantInsertPredicate, auth.tvfTenantUpdatePredicate',
        'H+': " COMPLETE AS OF PHASE 4 -- all ten objects exist and the build phase 1-4 is spent. auth.udfHasPermission "
              "(T-050) is two seeks with one asymmetry, the IsPlatformAdmin requirement of INV-09. auth.tvfPermissionScope "
              "(T-051) and the three predicates (T-061 to T-063) are INLINE TABLE-VALUED functions and are therefore "
              "named tvf and not udf: the house convention reserves udf for scalars, the SQL gate enforces it, and the "
              "design document was corrected rather than the code (BL-041). Every SESSION_CONTEXT read inside a "
              "predicate is TRY_CAST, because a predicate that throws fails the whole query instead of denying one row, "
              "and the permission ids are substituted as LITERALS by 120_rls_policy.sql.",
        'J': DEV,
        'K': 'Opus 5 (T-019, T-032, T-050, T-051, T-061, T-062, T-063, T-112)',
    },
    '170_permissions.sql': {
        'H+': " EXTENDED AT THE END OF PHASE 4, and both extensions were found by running it rather than by reading it. "
              "A new section 4 DENIES INSERT, UPDATE and DELETE on the four logs trail tables to applicationRole -- "
              "narrowing the conventions' SCHEMA::logs grant rather than revoking it, leaving SELECT alone and leaving "
              "logs.ExecutionLog alone because the framework writes it from the application side (BL-048); an "
              "impersonation probe confirmed a direct write now fails with 229 while the recorders still write through "
              "ownership chaining. The closing report now asserts ABSENCES as well as presences: ten Phase 3 and 4 "
              "objects that must be granted to nobody, each with its own reason, and the two session-context "
              "procedures that MUST carry applicationRole EXECUTE, because their absence would look exactly like empty "
              "data. The rlsBypassRole assertion was rewritten from 'holds nothing' to 'holds exactly the two "
              "maintenance EXECUTEs and nothing else' after the old form failed the deployment with severity 2 -- "
              "BL-047.",
        'E': '0, 2, 3, 4, 8',
        'J': DEV,
        'K': 'Opus 5 (G-19; BL-047, BL-048; T-104 lifts its assertion into 950)',
    },
    'Install-TemplateDatabase.ps1': {
        'H+': " AMENDED AGAIN THROUGH PHASES 2 TO 4. The manifest is now 28 steps and six of its orderings are "
              "load-bearing, listed in the file header: 025 ahead of every auth table because 110 reads its settings; "
              "165 ahead of 150, out of numeric order, because 150 THROWs at install time without "
              "logs.uspRecordAuthorizationDenial; 110 and 112 ahead of 105 because uspSetSessionContext calls "
              "auth.uspEndSession; 120 second to last; 170 last. A new Invoke-PolicyDrop step runs BEFORE every pass "
              "and calls auth.uspRebuildTenantAccessPolicy @Action = N'Drop', because a policy-bound SCHEMABINDING "
              "function cannot be altered (error 3729) and without the unbind every second deployment would fail at "
              "step 100. It asks master for DB_ID first so a first deployment is not aborted by -b, it honours -WhatIf, "
              "and its comment states plainly that the tables are unprotected between the drop and step 27. Full run "
              "on 2026-09-20: 28 steps, all OK, exit 0.",
        'J': DEV,
        'K': 'Opus 5 (T-011; amended T-022, and again for Phases 3 and 4)',
    },
}

# script name -> full replacement of D, H, I, J, K (the rows that were forward-looking 'Not started')
REPLACE = {
    '050_auth_permission.sql': dict(
        D='auth.PermissionCategory, auth.Permission',
        H="THE TABLES ARE HERE; THE 35 PERMISSIONS ARE NOT. Seeding DES Appendix A is T-089 in "
          "115_seed_reference_data.sql, and until it runs the catalogue is empty -- which is not harmless: "
          "120_rls_policy.sql resolves its permission ids from this table, finds none, substitutes the sentinel -1 and "
          "every predicate then denies every non-maintenance session. Fail-closed, reported by the script, and the "
          "reason 115 must re-run 120. PermissionDescription is an addition to the design (BL-037): without it the "
          "meaning of 35 permissions lives in the UI project. The category code is denormalised onto auth.Permission so "
          "the composite FK can keep the pair honest, the same arrangement as TenantTypeCode in 030.",
        I='Verified', J=DEV, K='Opus 5 (T-043; rows are T-089)'),
    '055_auth_role.sql': dict(
        D='auth.Role, auth.RolePermission',
        H="G-20 resolved: NO table-valued parameter types, and task T-045 was deleted rather than re-scoped. "
          "UX_auth_Role_Grant is filtered UNIQUE on (ApplicationId, OwnerTenantId, RoleCode) where IsDeleted = 0, so "
          "two tenants may both own a role called CaseWorker and a retired one does not block its replacement. INV-04 "
          "is structural here, not procedural: FK_auth_Role_OwnerTenant points at UX_auth_Tenant_Id_Application, the "
          "unfiltered pair added to 030 by ALTER (BL-040), and the script THROWs at install time if that constraint is "
          "missing rather than creating a table that cannot enforce its own invariant. The 14 baseline roles are seed "
          "data (T-089).",
        I='Verified', J=DEV, K='Opus 5 (T-044)'),
    '060_auth_profile_role.sql': dict(
        D='auth.UserProfileRole',
        H="The only table in this database that records a human decision, which is why it carries GrantedByProfileId "
          "and GrantedUtc as well as the grant itself. ScopeTenantId defaults to the profile's own tenant. "
          "CK_auth_UserProfileRole_Expiry requires ExpiresUtc > GrantedUtc -- so an already-expired grant has to be "
          "expressed by backdating GrantedUtc, which is what _tests/050 does. UX_auth_UserProfileRole_Grant is filtered "
          "UNIQUE on (UserProfileId, RoleId, ScopeTenantId) where IsDeleted = 0, and the consequence found by test is "
          "that a wholesale restore of soft-deleted grants collides at Msg 2601: a revoked grant is resurrected IN "
          "PLACE, one row at a time, never re-inserted.",
        I='Verified', J=DEV, K='Opus 5 (T-047)'),
    '065_auth_effective_permission.sql': dict(
        D='auth.ProfilePermissionScope, auth.uspRebuildProfilePermissionScope',
        H="The one piece of derived state in the design: flattened (UserProfileId, PermissionId, ScopeTenantId), never "
          "edited by hand, rebuilt whole for one profile or for the database. The rebuild excludes expired grants, "
          "deleted rows at every level and inactive or deleted users and profiles. "
          "IX_auth_ProfilePermissionScope_Lookup is FILTERED on IsDeleted = 0, which makes the filter load-bearing: a "
          "reader that omits IsDeleted = 0 from its own WHERE both misses the index and counts retired rows as live "
          "authority (BL-039). Proved by _tests/050 across fourteen mutations, compared against the set derived "
          "longhand from DES 8.6 in both directions after every one. The concurrency half of G-03 is still open.",
        I='Verified', J=DEV, K='Opus 5 (T-048, T-049; verified by _tests/050)'),
    '105_auth_session_procedures.sql': dict(
        D='auth.uspSetSessionContext, auth.uspClearSessionContext, auth.uspBeginMaintenanceSession, '
          'auth.uspEndMaintenanceSession',
        H="uspSetSessionContext takes @SessionTokenHash and NOT the token -- the design's sketch was corrected, not the "
          "code (BL-043) -- validates the whole chain (hash, expiry, user usable, profile usable, tenant usable, "
          "application match) and sets its five identity keys @read_only = 1. That is why uspClearSessionContext clears "
          "BypassRowSecurity and REPORTS the rest: a read-only key cannot be re-set, re-set to the same value, or "
          "nulled (Msg 15664), so one connection serves one profile (UI-06). A second call for the same profile returns "
          "silently; for a different profile it raises E-50022 and changes nothing. It runs AFTER 110 and 112 in the "
          "manifest because the expiry path calls auth.uspEndSession. The two maintenance procedures are gated on "
          "IS_ROLEMEMBER('rlsBypassRole') with E-50070 on refusal, demand a reason and a ticket, record ORIGINAL_LOGIN "
          "so the trail names the real login under EXECUTE AS, and write logs.AuthenticationEvent BEFORE setting the "
          "key so a bypass cannot exist without a trail. They are the only two objects granted to rlsBypassRole -- and "
          "that grant is what made 170's original assertion fail (BL-047).",
        I='Verified', J=DEV, K='Opus 5 (T-053, T-054, T-066; verified by _tests/050 and _tests/060)'),
    '120_rls_policy.sql': dict(
        D='auth.uspRebuildTenantAccessPolicy; auth.TenantAccessPolicy (FILTER plus three BLOCK predicates per '
          'registered table)',
        H="The policy is built by a PROCEDURE and not by inline DDL, and that is the decision this file exists to "
          "record (BL-038): a SCHEMABINDING function cannot be altered while a policy references it (error 3729), so a "
          "re-deployment needs a supported way to unbind first. auth.uspRebuildTenantAccessPolicy @Action = "
          "N'Rebuild'|N'Drop' reports @TablesBound and @TablesSkipped, and Install-TemplateDatabase.ps1 calls Drop "
          "before every pass and Rebuild at step 27. It reads config.TenantScopedTable, verifies each claimed TenantId "
          "column instead of trusting the registry, and substitutes the permission ids into the predicates as LITERALS "
          "-- a list, not a single id, because a database serving four applications holds four rows whose "
          "PermissionCode is Data.Read and one literal would protect one of them and fail open for three (BL-039). The "
          "closing report re-derives the list and CHARINDEXes it against the deployed definition, so a stale list is a "
          "failed deployment rather than a silent denial (UI-35). On dev: 2 tables bound, 8 predicates, state ON.",
        I='Verified', J=DEV, K='Opus 5 (T-064, T-065; verified by _tests/060)'),
    '150_auth_query_procedures.sql': dict(
        D='auth.uspDemandPermission (uspGetProfileContext, uspGetNavigationForProfile and the rest of the read surface '
          'are Phase 6 and are NOT in this file yet)',
        H="PARTIAL BY DESIGN and it says so when it runs: only auth.uspDemandPermission is Phase 3 work. It raises "
          "E-50030 when the profile holds the permission nowhere and E-50031 when it holds it but not at the acting "
          "tenant, two numbers because the UI's response differs (UI-13), and it writes logs.AuthorizationDenial "
          "through logs.uspRecordAuthorizationDenial before throwing -- which is why the manifest runs 165 BEFORE this "
          "file, out of numeric order, and why this file THROWs at install time if the recorder is absent. The limit is "
          "stated rather than papered over: inside a caller's open transaction the rollback takes the trail row with it "
          "(BL-042), so DES section 9 authorizes at step 5, before any transaction opens.",
        I='Verified', J=DEV, K='Opus 5 (T-052)'),
    '165_logs_procedures.sql': dict(
        D='logs.uspRecordAuthorizationChange, logs.uspRecordAuthorizationDenial, logs.uspRecordDataChange '
          '(logs.uspRecordAuthenticationEvent is written inline by 110 and 105 and is not a separate object here)',
        H="Never a session token and never a credential in the trail. All three recorders carry "
          "ActorAuthorityTenantId, so the record answers WHICH grant of authority was used and not merely who acted "
          "(P-08). EXECUTE is granted to NOBODY, deliberately: callers reach them through ownership chaining, which was "
          "measured on this instance and re-measured after 170 added its object-level DENY -- a direct UPDATE by an "
          "applicationRole member fails with 229 while the same write inside auth.uspRecordLoginFailure succeeds. "
          "Installed BEFORE 150 in the manifest because 150 asserts the denial recorder at install time.",
        I='Verified', J=DEV, K='Opus 5 (T-056)'),
}

NEW_ROWS = [
    ('database/_tests/050_authorization_and_session.sql', 'auth (test fixtures)',
     'No objects -- one AUTHZTEST application, a tenant tree, two profiles, three roles, five permissions and the '
     'grants the 14-mutation sequence edits',
     '3', '(040)', 'Yes',
     "NOT in the install manifest, for the same reason as 010 to 040: development fixtures with no business meaning. "
     "The two Phase 3 exit criteria that can only be settled by experiment. Section 4 is a loop over a 14-row "
     "#Mutation table -- grants, revocations, re-grants, a role edited while two profiles hold it, a permission "
     "retired out from under a role, an expiry, a profile deactivation and a user deactivation -- and after every one "
     "of them auth.ProfilePermissionScope is compared with the set derived longhand from DES 8.6, in BOTH directions "
     "and labelled by consequence: a missing row is a FALSE DENIAL, an extra row is a FALSE GRANT. The derived "
     "definition is written once, in one place, rather than reusing the procedure's own query, which would only have "
     "proved self-consistency. Section 5 MUST BE LAST: it spends the connection's identity, because the five keys are "
     "read-only once set. Three real constraints were found the hard way and are now expectations: "
     "CK_auth_UserProfileRole_Expiry, Msg 2601 on a wholesale restore of soft-deleted grants, and "
     "trg_au_updt_UserProfileRole overwriting auditDeletedBy so it cannot carry a caller's marker (BL-044). Last run: "
     "21 observations as intended, 3 notes, exit 0.",
     'Verified', DEV, 'Opus 5 (T-049, T-053, T-054, T-057, T-058)'),
    ('database/_tests/060_row_security.sql', 'auth, dbo (test fixtures)',
     'No objects that outlive the run -- a fixture across four tenants in two bound tables, plus one database user '
     'WITHOUT LOGIN in rlsBypassRole that section 8 drops again',
     '4', '(050)', 'Yes',
     "NOT in the install manifest. The Phase 4 exit criteria on live rows in two bound tables. It RE-RUNS "
     "auth.uspRebuildTenantAccessPolicy after building its fixture, because the predicates carry a literal permission-id "
     "list and a fixture that invents its own Data.Read row without rebuilding would test a policy that has never heard "
     "of its permissions -- every assertion would 'pass' by denying everything (UI-35). It plants rows with "
     "BypassRowSecurity set, then clears the key and CHECKS that it is clear before the matrix starts. Section 6 walks "
     "the whole DES 10.3 matrix and expects Msg 33504 by number, so a refusal for the wrong reason -- a check "
     "constraint, an FK, a permission -- fails the test. Two measured facts are written into the file: 33504 is the "
     "block-violation number, and the composite FKs on (AssignedToProfileId, TenantId) and (CaseFileId, TenantId) "
     "refuse a tenant move with Msg 547 BEFORE the predicate is consulted, which is why section 6b plants an "
     "unassigned row for 6g to move (BL-045). Section 5 shows a db_owner connection seeing 0 of 4 rows with no "
     "context, and section 7 does the same the supported way through a real rlsBypassRole member (UI-18). It found the "
     "CK_logs_AuthenticationEvent_Attributable defect in 105 (BL-046). Last run: 16 observations as intended, 3 notes, "
     "exit 0.",
     'Verified', DEV, 'Opus 5 (T-063, T-066, T-067, T-068)'),
]


def main():
    wb = load_workbook(WB)
    ws = wb['Scripts']
    col = {c: i + 1 for i, c in enumerate('ABCDEFGHIJK')}

    rowof = {}
    for r in ws.iter_rows(min_row=2):
        if r[1].value:
            rowof.setdefault(r[1].value, r[0].row)

    for name, changes in APPEND.items():
        rn = rowof[name]
        for key, val in changes.items():
            if key == 'H+':
                cur = ws.cell(row=rn, column=col['H']).value or ''
                ws.cell(row=rn, column=col['H']).value = cur.rstrip() + val
            else:
                ws.cell(row=rn, column=col[key]).value = val

    for name, changes in REPLACE.items():
        rn = rowof[name]
        for key, val in changes.items():
            ws.cell(row=rn, column=col[key]).value = val

    # --- two new test rows, inserted after the last existing artefact row -----------------------------
    tools_row = rowof['tools/T040-PasswordVerification (.NET 10 console)']
    tpl = rowof['database/_tests/040_identity_and_authn.sql']
    marker = ws.cell(row=tpl, column=1).value           # the same glyph the other test rows use
    at = tools_row + 1
    ws.insert_rows(at, amount=len(NEW_ROWS))
    for i, vals in enumerate(NEW_ROWS):
        rn = at + i
        ws.cell(row=rn, column=1).value = marker
        for j, v in enumerate(vals):
            ws.cell(row=rn, column=j + 2).value = v
        for c in range(1, 12):
            ws.cell(row=rn, column=c)._style = copy.copy(ws.cell(row=tpl, column=c)._style)
        ws.row_dimensions[rn].height = ws.row_dimensions[tpl].height

    # --- extend the auto-filter over the two new rows -------------------------------------------------
    if ws.auto_filter.ref:
        head, tail = ws.auto_filter.ref.split(':')
        lastcol = ''.join(ch for ch in tail if ch.isalpha())
        lastrow = int(''.join(ch for ch in tail if ch.isdigit())) + len(NEW_ROWS)
        ws.auto_filter.ref = '%s:%s%d' % (head, lastcol, lastrow)

    wb.save(WB)
    print('appended', len(APPEND), 'replaced', len(REPLACE), 'inserted at', at, 'filter', ws.auto_filter.ref)


if __name__ == '__main__':
    main()
