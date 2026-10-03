# Records the four template defects found by the WellDrillersLicense build and fixed here on 2026-09-25:
# G-52 .. G-55 in gaps.xlsx, BL-086 in the build log, and the new test file in the Scripts sheet.
import copy
from openpyxl import load_workbook

GAPS = r'../workbooks/gaps.xlsx'
BUILD = r'../workbooks/build-and-traceability.xlsx'

G = [
    ['G-52', 'A refused sign-in was enrolment proof for ANY account, so the password alone could issue recovery codes to an account that already had a factor',
     'DES 7.x MFA enrolment; T-041; E-50122',
     'auth.udfResolveEnrolmentActor is a route to a FIRST factor and nothing else: a refused MfaRequired exchange proves the password, never possession of a factor.',
     'The function checked only the exchange (Failure, MfaRequired, PasswordVerified, within the window). It left the first-factor rule to its callers, and they did not hold it: auth.uspEnrolMfaFactor refuses only a confirmed factor OF THE SAME TYPE (E-50119), and auth.uspIssueMfaRecoveryCodes REQUIRES a confirmed factor.',
     'Account takeover with the password alone. For an account with a confirmed TOTP factor, sign in with the password, get refused E-50109, and spend that refusal on a batch of recovery codes: one-time passwords that satisfy the MFA the attacker does not have. It also replaces the owner\'s own codes. Reproduced on a clean build: frank\'s 4 codes became the attacker\'s 1.',
     'Critical', 'Any deployment with RequireMfaForLocal = 1',
     'The rule lives in the function, once: return NULL when the account held a confirmed factor at or before AttemptedUtc. Factors confirmed after the exchange do not count, so the legitimate enrol, confirm, recovery-codes visit still works.',
     'None', '2 Identity', 'Closed',
     'CLOSED 2026-09-25 in 100_auth_functions.sql. Proved by _tests/040 section 12m\', which failed before the fix (codes issued) and passes after it (E-50122, no codes). Found by the WellDrillersLicense build (its GAP-WDL-005 / BL-WDL-023), where WebAuthn credentials count as well.'],
    ['G-53', 'E-50061 and E-50066 are reachable and never raised by any test, so 080 fails its own coverage check on a clean build',
     'DES Appendix B; 080 section 14',
     'Every number a shipped module can throw is raised by some test, or is on 080\'s list of numbers that cannot be.',
     'No file under _tests calls auth.uspRegisterExternalUser on a registration that is not Approved (E-50061), or with a user name in use (E-50066).',
     'The regression suite could not go green on a fresh database, so a green run meant a database that had been hand-nursed.',
     'Medium', 'A clean-build regression run',
     'Add _tests/075_registration_probes.sql, which provokes both. It runs between 070, which leaves an Approved registration, and 080.',
     'None', '8 Verification', 'Closed',
     'CLOSED 2026-09-25 by _tests/075_registration_probes.sql. Found by the WellDrillersLicense build (its BL-WDL-003).'],
    ['G-54', '900 wrote a root-policy PolicyNote longer than NVARCHAR (1000), so the bootstrap failed with Msg 2628 on every fresh database',
     'DES 16.2; T-092',
     'The first administrator can be bootstrapped on a clean database.',
     'The note literal in 900_bootstrap_first_admin.sql had grown past the column.',
     'No fresh install could bootstrap, so 070 and 080 could never run on a clean build. The existing test databases predated the longer note.',
     'High', 'Every fresh install',
     'Shorten the note, keeping its instruction (set RequireMfaForLocal and RequireStepUpForPrivileged to 1 once the administrator has a factor; G-30, G-37).',
     'None', '8 Verification', 'Closed',
     'CLOSED 2026-09-25 in 900_bootstrap_first_admin.sql. Found by the WellDrillersLicense build (its BL-WDL-001).'],
    ['G-55', 'Re-granting a role after a revoke through auth.uspAssignRoleToProfile always failed with E-50010',
     'DES 8.6; T-047; P-07',
     'A revoked grant comes back to life in place when it is granted again, so its history stays one row.',
     'The resurrect branch of auth.uspAssignRoleToProfile re-stamps GrantedUtc and GrantedByProfileId, as it should, since it is a new decision by a new person. trg_au_updt_UserProfileRole refused any change to either column.',
     'Any revoke-then-regrant failed. _tests/050 resurrects with a direct UPDATE that leaves both columns alone, so the suite never saw it.',
     'High', 'Role administration: every re-grant',
     'The trigger allows GrantedUtc and GrantedByProfileId to change on a resurrection (IsDeleted 1 to 0 in the same statement) and nowhere else. UserProfileId, RoleId, ScopeTenantId and ApplicationId stay immutable on every row.',
     'None', '3 Authorization', 'Closed',
     'CLOSED 2026-09-25 in 060_auth_profile_role.sql. Proved by _tests/080 section 8b\'\', which failed with E-50010 before the fix and passes after it (the same row resurrected). Found by the WellDrillersLicense four-eyes test (its BL-WDL-025).'],
]

BL = ['BL-086', '2026-09-25',
      'database/060_auth_profile_role.sql, database/100_auth_functions.sql, database/900_bootstrap_first_admin.sql, database/_tests/040_identity_and_authn.sql, database/_tests/075_registration_probes.sql, database/_tests/080_error_catalogue.sql, database/_tests/Build-Upstream.sh',
      'DES 7.x, 8.6, 16.2, Appendix B; G-52, G-53, G-54, G-55',
      'The template installs, bootstraps and passes 010-080 on a clean database. A revoked grant can be granted again. Enrolment proof is first-factor only.',
      'Four defects found downstream by the WellDrillersLicense build and carried back. For each, the test was written first, run on an unfixed clean build (tplUpstream) where it failed, then passed after the fix. G-52: udfResolveEnrolmentActor returns NULL for an account with a factor confirmed before the exchange; 040 12m\'. G-53: 075_registration_probes.sql raises E-50061 and E-50066. G-54: the 900 PolicyNote fits NVARCHAR (1000). G-55: the UserProfileRole trigger allows a resurrection to re-stamp GrantedUtc and GrantedByProfileId; 080 8b\'\'. Build-Upstream.sh builds a disposable database and runs the suite.',
      'Nothing in the design changes: each fix makes the code do what the design already said. G-52 is the serious one, a password-only route to working recovery codes. It is closed in the function rather than in each caller, because the callers are where it was lost.',
      'Build, carrying back downstream findings.',
      'No design change needed. The design already required each behaviour.']

SCRIPT = ['—','database/_tests/075_registration_probes.sql', 'auth (test probes)',
          'No objects: one Pending registration, REG075PEND, created once and found again on a re-run',
          '6', '(070)', 'Yes',
          'NOT in the install manifest. Provokes the two auth.uspRegisterExternalUser refusals no other test reaches (E-50061, E-50066), so 080\'s coverage check passes on a clean build (G-53).',
          'Verified', '2026-09-25  MDE-55TT2J4\\tplUpstream', 'Opus 5.5 (G-53)']


def append(ws, rows):
    tpl = ws.max_row
    for row in rows:
        rn = ws.max_row + 1
        for c, v in enumerate(row, start=1):
            ws.cell(row=rn, column=c, value=v)
            ws.cell(row=rn, column=c)._style = copy.copy(ws.cell(row=tpl, column=c)._style)


wb = load_workbook(GAPS)
ws = wb['Gaps']
assert not any(r[0] in ('G-52', 'G-53', 'G-54', 'G-55') for r in ws.iter_rows(min_row=2, values_only=True)), 'already recorded'
append(ws, G)
wb.save(GAPS)

wb = load_workbook(BUILD)
ws = wb['Build Log']
assert not any(r[0] == 'BL-086' for r in ws.iter_rows(min_row=2, values_only=True)), 'already recorded'
append(ws, [BL])
ws = wb['Scripts']
at = next(i for i, r in enumerate(ws.iter_rows(min_row=1, values_only=True), start=1)
          if r[1] == 'database/_tests/080_error_catalogue.sql')
ws.insert_rows(at, 1)
for c, v in enumerate(SCRIPT, start=1):
    ws.cell(row=at, column=c, value=v)
    ws.cell(row=at, column=c)._style = copy.copy(ws.cell(row=at + 1, column=c)._style)
wb.save(BUILD)
print('workbooks updated')
