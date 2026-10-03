import sys
from openpyxl import load_workbook

wb = load_workbook('workbooks/implementation-tracking.xlsx', data_only=True)
ws = wb['Tasks']
hdr = [c.value for c in ws[1]]
want = set(sys.argv[1:]) if len(sys.argv) > 1 else None
idx = {h: i for i, h in enumerate(hdr)}
for row in ws.iter_rows(min_row=2, values_only=True):
    tid = row[idx['Task ID']]
    if not tid:
        continue
    n = int(str(tid).split('-')[1])
    if want and str(n).zfill(3) not in want and tid not in want:
        continue
    print('=' * 110)
    print(f"{tid} | Phase {row[idx['Phase']]} | {row[idx['Status']]} | {row[idx['% Done']]}% | Est {row[idx['Est. Days']]}")
    print(f"  Artefact : {row[idx['Artefact / Script']]}")
    print(f"  Task     : {row[idx['Task']]}")
    print(f"  Type     : {row[idx['Type']]}   Design Ref: {row[idx['Design Ref']]}   Depends: {row[idx['Depends On']]}")
    print(f"  Notes    : {row[idx['Notes / Blockers']]}")
