"""Bring the Traceability sheet up to date for Phases 3 and 4."""
import copy
from openpyxl import load_workbook

WB = 'workbooks/build-and-traceability.xlsx'

# row -> {column letter: value}
EDITS = {
    6: {'E': '_tests/060_row_security.sql section 6a (a profile scoped above sees 3 of 4 rows)', 'F': 'Partial',
        'G': 'The MECHANISM is verified: a read grant at an ancestor tenant makes the descendants visible through the '
             'FILTER predicate. What is not verified is this application\'s seeding -- which role holds Data.Read at '
             'the root -- which is T-089.'},
    7: {'E': '_tests/060_row_security.sql sections 6a and 6f', 'F': 'Partial',
        'G': 'Verified as a predicate outcome: a row outside the profile\'s scope is invisible to a read and refused '
             'to an update with Msg 33504. INV-05 in uspAssignRoleToProfile (145) is Phase 6.'},
    11: {'D': 'auth.tvfTenantInsertPredicate (100); auth.TenantAccessPolicy (120)',
         'E': '_tests/060_row_security.sql sections 6b and 6c (insert into the acting tenant succeeds; into a sibling '
              'and into a descendant both 33504)', 'F': 'Verified',
         'G': 'The insert predicate demands EQUALITY with the acting tenant, not descendant-of, which is what makes an '
              'agency user\'s entry land in the agency\'s own tenant and stay invisible to the county. UI-14: the UI '
              'never offers a TenantId picker on an insert.'},
    24: {'D': 'Read grant at the root tenant; auth.tvfTenantReadPredicate (100), policy (120)',
         'E': '_tests/060_row_security.sql section 6a', 'F': 'Partial',
         'G': 'Same as App 1: the predicate is verified, the seeding that puts MDE staff at the root is T-089.'},

    40: {'D': '050_auth_permission.sql (auth.PermissionCategory, auth.Permission, PermissionDescription); '
              '115_seed_reference_data.sql seeds the 35 rows and has NOT run',
         'E': '050 closing report; _tests/050 and _tests/060 run against their own fixture permissions',
         'F': 'Partial',
         'G': 'The tables, the composite FK on (PermissionCategoryId, PermissionCategoryCode) and the description '
              'column are built and exercised. The catalogue itself is empty until T-089, and that is NOT neutral: '
              '120_rls_policy.sql resolves its permission ids from this table, finds none, substitutes the sentinel -1 '
              'and every predicate then denies every non-maintenance session. Fail-closed and reported, but a project '
              'that deploys without seeding will see empty screens (UI-35).'},
    41: {'D': '055_auth_role.sql; FK_auth_Role_OwnerTenant -> UX_auth_Tenant_Id_Application (030, BL-040); INV-04 also '
              'checked in uspAssignRoleToProfile (145, Phase 6)',
         'E': '055 closing report; _tests/050 fixture builds three roles across two tenants', 'F': 'Built',
         'G': 'The half of INV-04 that says a role and its owner tenant belong to one application is now STRUCTURAL -- '
              'a composite foreign key, not a procedure check, so it holds for the seed script and for SSMS too. The '
              'grant-side half is 145 and remains T-091.'},
    42: {'D': '060_auth_profile_role.sql (ScopeTenantId, GrantedByProfileId, ExpiresUtc); '
              '065_auth_effective_permission.sql flattens them',
         'E': '_tests/050_authorization_and_session.sql sections 3 and 4', 'F': 'Verified',
         'G': 'D-03 holds: a grant carries its own scope, so one profile can hold different permissions at different '
              'tenants. CK_auth_UserProfileRole_Expiry (ExpiresUtc > GrantedUtc) and the filtered uniqueness of a live '
              'grant are both exercised by the mutation loop.'},
    43: {'D': '065_auth_effective_permission.sql: auth.ProfilePermissionScope, auth.uspRebuildProfilePermissionScope',
         'E': '_tests/050_authorization_and_session.sql section 4 -- 14 mutations, compared against the set derived '
              'longhand from DES 8.6 in BOTH directions after each one, 0 failed', 'F': 'Verified',
         'G': 'The strongest verified row on this sheet, and deliberately so: the materialization is the one piece of '
              'derived state in the design, and a wrong answer is either a false denial (an empty screen) or a false '
              'grant (nobody complains). The comparison is written longhand rather than reusing the procedure\'s query, '
              'so it can disagree with it. The concurrency half of G-03 -- two rebuilds for one profile at once -- is '
              'still open.'},
    44: {'D': '105_auth_session_procedures.sql (steps 1-4); 150_auth_query_procedures.sql uspDemandPermission (step 5); '
              '100/120 predicates (step 6)',
         'E': '_tests/050 sections 4b and 5; _tests/060 sections 5 to 7', 'F': 'Partial',
         'G': 'Every step of the path now exists and each has been exercised. What is not verified is the path being '
              'FOLLOWED: the domain procedures that must call uspDemandPermission at step 5, before opening a '
              'transaction, are Phase 7 (180) -- and BL-042 is the reason the ordering is a requirement and not a '
              'style preference.'},
    45: {'D': '025_config_tables.sql config.TenantScopedTable (2 live rows); auth.uspRebuildTenantAccessPolicy (120) '
              'builds the policy from it',
         'E': '120 closing report (2 table(s) protected, 0 registry row(s) skipped); _tests/060 section 4 re-runs the '
              'rebuild and prints what was bound', 'F': 'Verified',
         'G': 'The registry is the only place the protected-table list exists, so protecting a new table is an INSERT '
              'plus a rebuild rather than an edit to a CREATE SECURITY POLICY statement. A registry row is treated as a '
              'claim, not a fact: the rebuild verifies the TenantId column and reports @TablesSkipped. The 950 '
              'assertion that no registered table is left unprotected is still Phase 8 (INV-13).'},
    46: {'D': '100_auth_functions.sql: auth.tvfTenantReadPredicate, auth.tvfTenantInsertPredicate, '
              'auth.tvfTenantUpdatePredicate (WITH SCHEMABINDING)',
         'E': '120 closing report re-derives the permission-id list and CHARINDEXes it against the deployed '
              'definitions; _tests/060 section 6 exercises all three', 'F': 'Verified',
         'G': 'Two corrections to the design landed here: the predicates carry a LITERAL LIST of permission ids because '
              'a schema-bound function cannot join auth.Permission by code and because four applications hold four '
              'Data.Read rows (BL-039), and all four table-valued functions are named tvf, not udf (BL-041). Every '
              'SESSION_CONTEXT read is TRY_CAST: a predicate that throws fails the query instead of denying the row.'},
    47: {'D': '100_auth_functions.sql three predicates; auth.TenantAccessPolicy (120) binds FILTER once and BLOCK '
              'three times per table',
         'E': '_tests/060_row_security.sql section 6 -- the full DES 10.3 matrix, eight cells, each refusal expected by '
              'number (Msg 33504)', 'F': 'Verified',
         'G': 'The asymmetry is the point and the test proves both halves: a row outside scope is INVISIBLE to a read '
              'and LOUD on a write. One wrinkle is recorded as BL-045 and belongs on this row: the composite foreign '
              'keys on the demo tables refuse a tenant move with Msg 547 BEFORE the block predicate is consulted, so a '
              'UI mapping errors to messages must handle 547 on those tables as "this row cannot be moved".'},
    48: {'D': 'auth.uspDemandPermission (150); every domain procedure must call it -- 180 is Phase 7',
         'E': '_tests/050 exercises the procedure; the CALLERS are not written', 'F': 'Partial',
         'G': 'The gap is unchanged and is G-05: RLS filters rows, it does not know which VERB is being attempted, so '
              'a profile with read-only authority can still be handed an update procedure and be refused only by the '
              'block predicate -- with the wrong error number for a UI. uspDemandPermission is the intended guard and '
              'now exists; nothing yet forces a procedure to call it.'},
    49: {'D': 'auth.uspBeginMaintenanceSession, auth.uspEndMaintenanceSession (105); rlsBypassRole (005); '
              'BypassRowSecurity session key honoured by all three predicates',
         'E': '_tests/060_row_security.sql sections 5 and 7 -- 4 rows inside the window, 0 after the end call, '
              '@BypassCleared = 1, and both logs.AuthenticationEvent rows naming the real login', 'F': 'Verified',
         'G': 'Verified end to end through a real member of rlsBypassRole (a user WITHOUT LOGIN, created and dropped by '
              'the test), which is what it took to reach the accepted path at all -- the deployment probe inside 105 '
              'can only reach the refusal path, and that is how the CK_logs_AuthenticationEvent_Attributable defect of '
              'BL-046 survived until Phase 4. Section 5 also records the UI-18 fact: a db_owner connection with no '
              'session context sees ZERO rows, which looks exactly like data loss.'},
    55: {'D': '105_auth_session_procedures.sql, 150 and 165 follow the full contract; every procedure\'s preamble',
         'E': '_tests/050 and _tests/060; the SQL gate checks the Rule 8 TRY/CATCH shape on every file', 'F': 'Partial',
         'G': 'The contract holds for every procedure written so far, and the gate enforces the mechanical half of it. '
              'Partial because two thirds of the procedures in the design are still unwritten.'},
    63: {'D': '040_auth_userprofile.sql UX_auth_UserProfile_Default, filtered on IsDefault = 1 AND IsDeleted = 0',
         'E': '040 closing report; _tests/050 builds two profiles for one user', 'F': 'Built',
         'G': 'The index makes a second default unrepresentable rather than merely forbidden. The negative case -- a '
              'procedure trying to set a second default and being refused cleanly rather than by a constraint error -- '
              'is uspCreateProfile in 140 and remains T-091.'},
    64: {'D': '055_auth_role.sql composite FK (BL-040) for the owner half; 145_auth_role_procedures.sql for the grant '
              'half', 'E': '055 install-time assertion; T-091 for the grant half', 'F': 'Partial',
         'G': 'Half of this invariant became structural in Phase 3 and no longer depends on a procedure being used.'},
    71: {'G': 'Verified by the closing report and by an impersonation probe. The ADJACENT hole was found during the '
              'Phase 4 closeout and closed in the same file: the conventions\' permission script grants all four table '
              'verbs on SCHEMA::logs, so the audit trails were directly writable even though SCHEMA::auth was denied '
              '-- 170 now denies INSERT, UPDATE and DELETE on the four trail tables object by object (BL-048). INV-11 '
              'was never the whole of the "no direct table access" rule; it was the half that had been written down.'},
    72: {'D': 'auth.tvfTenantUpdatePredicate (100) and the BLOCK AFTER UPDATE binding (120); composite FKs on the demo '
              'tables; 135_audit_triggers.sql (Phase 7) for the column-level refusal',
         'E': '_tests/060_row_security.sql section 6g -- the tenant move is refused', 'F': 'Partial',
         'G': 'The move is refused today, but by TWO mechanisms and not the one the design names: the composite '
              'foreign key fires first with Msg 547, and the block predicate catches what it does not with Msg 33504 '
              '(BL-045). Neither produces the clean, named error a UI wants; that is the trigger in 135.'},
    73: {'D': '120_rls_policy.sql closing report (every registry row either bound or reported skipped); '
              '950_verify_deployment.sql is the standing check',
         'E': '120 closing report on every deployment: 2 of 2 bound, 0 skipped', 'F': 'Partial',
         'G': 'Checked at deployment time already, which is the moment it can still be fixed. Partial because nothing '
              'checks it BETWEEN deployments -- a registry row inserted by an administrator without a rebuild leaves a '
              'table unprotected and silent, and that standing check is 950.'},
}

NEW = [
    ('Design', 'DES 14.3 One connection, one profile',
     'A second auth.uspSetSessionContext on the same connection returns silently for the same profile and refuses a '
     'different one; the refusal changes nothing.',
     'auth.uspSetSessionContext sets its five identity keys with @read_only = 1 (105); auth.uspClearSessionContext '
     'clears only BypassRowSecurity',
     '_tests/050_authorization_and_session.sql section 5 -- same profile silent with UserProfileId intact, different '
     'profile Error 50022, identity keys survive the refusal',
     'Verified',
     'This is the row behind UI-06, and the mechanism is stronger than the contract: a read-only session key cannot be '
     're-set to a different value, re-set to the SAME value, or nulled (Msg 15664, measured here). So the refusal is '
     'not a policy the procedure could relax -- one connection serves one profile for as long as it stays open, and a '
     'connection pool that hands a pooled connection to a second user is a defect the database will catch.'),
    ('Design', 'DES 21.3 Schema-bound predicates and re-deployment',
     'A function referenced by a security policy cannot be altered; a deployment must therefore unbind before it '
     'rebuilds.',
     'auth.uspRebuildTenantAccessPolicy @Action = N\'Drop\' | N\'Rebuild\' (120); Invoke-PolicyDrop before every pass '
     'in Install-TemplateDatabase.ps1 (step 0), rebuild at step 27',
     'Two full 28-step runs on 2026-09-20 against an existing database: the unbind reported OK, step 21 re-created the '
     'functions, step 27 reported 2 table(s) protected and 0 registry row(s) skipped, exit 0',
     'Verified',
     'The design filed this as a caution; it is a deployment blocker (error 3729), and every second deployment of an '
     'existing database would have failed without the unbind -- BL-038. The cost is stated rather than hidden: between '
     'the drop and step 27 the tenant-scoped tables are UNPROTECTED, which is why the drop is a deployment step on a '
     'maintenance connection and never an operational one.'),
    ('Design', 'DES 10.2 Behaviour with an unseeded catalogue',
     'Not in the design, and it should have been: what the predicates do before 115_seed_reference_data.sql has ever '
     'run.',
     '120_rls_policy.sql substitutes the sentinel PermissionId IN (-1) when the catalogue resolves to no rows, and '
     'reports it',
     '120 closing report on the current dev database; _tests/060 section 4 prints the lists that ended up in the '
     'functions',
     'Verified',
     'It fails CLOSED -- every non-maintenance session sees nothing -- which is the right direction and the wrong '
     'experience, so it is written down in three places rather than left to be discovered: here, in the script\'s own '
     'report, and as UI-35. The same shape recurs whenever the catalogue changes without a rebuild, and that case is '
     'worse, because it is partial rather than total.'),
]


def main():
    wb = load_workbook(WB)
    ws = wb['Traceability']
    col = {c: i + 1 for i, c in enumerate('ABCDEFG')}

    for rn, changes in EDITS.items():
        for k, v in changes.items():
            ws.cell(row=rn, column=col[k]).value = v

    tpl = 76
    at = 77
    ws.insert_rows(at, amount=len(NEW))
    for i, vals in enumerate(NEW):
        rn = at + i
        for j, v in enumerate(vals):
            ws.cell(row=rn, column=j + 1).value = v
        for c in range(1, 8):
            ws.cell(row=rn, column=c)._style = copy.copy(ws.cell(row=tpl, column=c)._style)
        ws.row_dimensions[rn].height = ws.row_dimensions[tpl].height

    if ws.auto_filter.ref:
        head, tail = ws.auto_filter.ref.split(':')
        lastcol = ''.join(ch for ch in tail if ch.isalpha())
        lastrow = int(''.join(ch for ch in tail if ch.isdigit())) + len(NEW)
        ws.auto_filter.ref = '%s:%s%d' % (head, lastcol, lastrow)

    wb.save(WB)
    print('traceability edits', len(EDITS), 'new rows', len(NEW), 'filter', ws.auto_filter.ref)


if __name__ == '__main__':
    main()
